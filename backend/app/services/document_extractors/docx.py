from io import BytesIO

from docx import Document


def extract_document(content: bytes) -> tuple[str | None, str]:
    document = Document(BytesIO(content))
    paragraphs = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    table_rows = [
        "\t".join(cell.text.strip() for cell in row.cells)
        for table in document.tables
        for row in table.rows
        if any(cell.text.strip() for cell in row.cells)
    ]

    extracted_text = "\n".join([*paragraphs, *table_rows]).strip()
    return extracted_text, "extracted"
